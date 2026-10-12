# NCUACG/songfest/importers.py
"""
從社團的 Google 試算表（下載成 .xlsx）同步舊曲與播放紀錄。

試算表是正本：每次匯入都把網站資料改成跟試算表一樣，可以重複執行。
  - 線上歌單：以「DVD編號」分頁為主資料；其他五十音分頁只用來決定分類與排列順序。
  - 播放紀錄：每個播放日期是一場歌謠祭；該場的播放紀錄會整批換成試算表裡的內容。

後台的「匯入試算表」頁面與 `python manage.py import_songfest` 都呼叫這裡。
"""
from __future__ import annotations

import datetime as dt
import re
import warnings
from collections import defaultdict
from dataclasses import dataclass, field

import openpyxl
from django.db import transaction

from .models import Festival, Play, Track
from .text import cell_text, duration_seconds, guess_festival_name, normalize, play_date, song_key

MAIN_SHEET = "DVD編號"
_DVD_CODE_RE = re.compile(r"^D\d+$")

# 五十音的一行：あ行包含あいうえお，篩選「あ行」時要一起列出
_KANA_ROWS = ["あいうえお", "かきくけこ", "さしすせそ", "たちつてと", "なにぬねの",
              "はひふへほ", "まみむめも", "やゆよ", "らりるれろ", "わをん"]
KANA_GROUP = {kana: row[0] for row in _KANA_ROWS for kana in row}

TRACK_FIELDS = [
    "song_source", "visual_source", "title", "artist", "editor", "duration_seconds",
    "index_label", "index_group", "index_order",
]
DERIVED_FIELDS = ["dvd_sort", "song_key", "search_text"]


class SpreadsheetError(ValueError):
    """試算表格式不符（找不到需要的分頁或欄位），訊息會直接顯示給使用者。"""


@dataclass
class ImportReport:
    lines: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def add(self, line: str) -> None:
        self.lines.append(line)

    def warn(self, line: str) -> None:
        self.warnings.append(line)


@dataclass
class SongRow:
    dvd_code: str
    song_source: str
    visual_source: str
    title: str
    artist: str
    editor: str
    duration_seconds: int | None
    index_label: str = ""
    index_group: str = ""
    index_order: int = 0


@dataclass
class PlayRow:
    date: dt.date
    song_source: str
    visual_source: str
    title: str


# ── 讀試算表 ────────────────────────────────────────────────────────────────

def load_workbook(source):
    """source 可以是檔案路徑或上傳的檔案物件。"""
    with warnings.catch_warnings():
        # 播放日期欄被設成日期格式但內容是 20241214 這種數字時，openpyxl 會警告；我們自己會轉
        warnings.simplefilter("ignore", UserWarning)
        try:
            return openpyxl.load_workbook(source, read_only=True, data_only=True)
        except Exception as exc:  # 檔案壞掉、不是 xlsx 等
            raise SpreadsheetError(f"無法讀取試算表，請確認是 .xlsx 檔：{exc}") from exc


def _rows(sheet) -> list[tuple]:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        return list(sheet.iter_rows(values_only=True))


def _find_header(rows, required: set[str], search_rows: int = 10):
    """在前幾列找出標題列，回傳 (標題列位置, {欄名: 欄位索引})；找不到回傳 None。"""
    for i, row in enumerate(rows[:search_rows]):
        columns = {}
        for j, value in enumerate(row):
            name = cell_text(value)
            if name and name not in columns:
                columns[name] = j
        if required <= columns.keys():
            return i, columns
    return None


def _get(row, columns, name):
    j = columns.get(name)
    return row[j] if j is not None and j < len(row) else None


def read_songlist(workbook, report: ImportReport) -> list[SongRow]:
    required = {"DVD編號", "歌曲出處", "曲名"}
    main = None
    if MAIN_SHEET in workbook.sheetnames:
        main = workbook[MAIN_SHEET]
    else:
        main = next((ws for ws in workbook.worksheets if _find_header(_rows(ws), required | {"長度"})), None)
    if main is None:
        raise SpreadsheetError("線上歌單裡找不到「DVD編號」分頁（需要 DVD編號、歌曲出處、曲名、長度 欄位）")

    rows = _rows(main)
    header = _find_header(rows, required)
    if header is None:
        raise SpreadsheetError(f"「{main.title}」分頁找不到 DVD編號、歌曲出處、曲名 的標題列")
    start, columns = header

    songs: dict[str, SongRow] = {}
    for row in rows[start + 1:]:
        code = cell_text(_get(row, columns, "DVD編號"))
        if not code:
            continue
        if not _DVD_CODE_RE.match(code):
            report.warn(f"DVD 編號格式不對，已略過：{code}")
            continue
        if code in songs:
            report.warn(f"DVD 編號重複，只保留第一筆：{code}")
            continue
        title = cell_text(_get(row, columns, "曲名"))
        source = cell_text(_get(row, columns, "歌曲出處"))
        if not title or not source:
            report.warn(f"{code} 缺少歌曲出處或曲名，已略過")
            continue
        songs[code] = SongRow(
            dvd_code=code,
            song_source=source,
            visual_source=cell_text(_get(row, columns, "畫面出處")),
            title=title,
            artist=cell_text(_get(row, columns, "歌手")),
            editor=cell_text(_get(row, columns, "製作者")),
            duration_seconds=duration_seconds(_get(row, columns, "長度")),
        )

    # 五十音分頁：小節標題列（例如「あ 行」）只有第一欄有字
    order = 0
    indexed = set()
    for sheet in workbook.worksheets:
        if sheet.title == main.title or "封存" in sheet.title:
            continue
        rows = _rows(sheet)
        header = _find_header(rows, required)
        if header is None:
            continue
        start, columns = header
        label = sheet.title
        for row in rows[start + 1:]:
            code = cell_text(_get(row, columns, "DVD編號"))
            source = cell_text(_get(row, columns, "歌曲出處"))
            if not code:
                if source and not cell_text(_get(row, columns, "曲名")):
                    label = source.replace("行", "").strip()
                continue
            song = songs.get(code)
            if song is None or code in indexed:
                continue
            order += 1
            indexed.add(code)
            song.index_label = label
            song.index_group = KANA_GROUP.get(label, label)
            song.index_order = order

    missing = [code for code in songs if code not in indexed]
    for code in missing:
        # 排在最後面，依 DVD 編號
        songs[code].index_order = 1_000_000 + int(code[1:])
    if missing:
        report.warn(f"{len(missing)} 首沒有出現在任何五十音分頁，會排在清單最後：{'、'.join(missing[:10])}")

    return list(songs.values())


def read_play_history(workbook, report: ImportReport) -> list[PlayRow]:
    required = {"歌曲出處", "曲名", "播放日期"}
    for sheet in workbook.worksheets:
        rows = _rows(sheet)
        header = _find_header(rows, required)
        if header is not None:
            break
    else:
        raise SpreadsheetError("播放紀錄裡找不到有 歌曲出處、曲名、播放日期 欄位的分頁")

    start, columns = header
    plays: list[PlayRow] = []
    bad_dates = 0
    for row in rows[start + 1:]:
        title = cell_text(_get(row, columns, "曲名"))
        source = cell_text(_get(row, columns, "歌曲出處"))
        if not title and not source:
            continue  # 空白列
        date = play_date(_get(row, columns, "播放日期"))
        if date is None or not title:
            bad_dates += 1
            continue
        plays.append(PlayRow(date, source, cell_text(_get(row, columns, "畫面出處")), title))
    if bad_dates:
        report.warn(f"{bad_dates} 列缺少曲名或播放日期看不懂，已略過")
    return plays


# ── 寫入資料庫 ──────────────────────────────────────────────────────────────

def sync_tracks(songs: list[SongRow], report: ImportReport, prune: bool = False) -> None:
    existing = {t.dvd_code: t for t in Track.objects.all()}
    to_create, to_update = [], []
    for song in songs:
        values = {name: getattr(song, name) for name in TRACK_FIELDS}
        track = existing.get(song.dvd_code)
        if track is None:
            track = Track(dvd_code=song.dvd_code, **values)
            track.fill_derived()
            to_create.append(track)
        elif any(getattr(track, name) != value for name, value in values.items()):
            for name, value in values.items():
                setattr(track, name, value)
            track.fill_derived()
            to_update.append(track)

    Track.objects.bulk_create(to_create, batch_size=500)
    Track.objects.bulk_update(to_update, TRACK_FIELDS + DERIVED_FIELDS, batch_size=200)

    unchanged = len(songs) - len(to_create) - len(to_update)
    report.add(f"舊曲：新增 {len(to_create)}、更新 {len(to_update)}、沒有變動 {unchanged}")

    missing = sorted(set(existing) - {s.dvd_code for s in songs})
    if missing:
        if prune:
            Track.objects.filter(dvd_code__in=missing).delete()
            report.add(f"已刪除試算表裡沒有的 {len(missing)} 首：{'、'.join(missing[:10])}")
        else:
            report.warn(
                f"網站上有 {len(missing)} 首試算表裡沒有（沒有刪除，要刪請勾選「刪除試算表沒有的舊曲」）："
                + "、".join(missing[:10])
            )


def sync_plays(plays: list[PlayRow], report: ImportReport) -> None:
    # 播放紀錄沒有 DVD 編號，用正規化後的「歌曲出處 + 畫面出處 + 曲名」對歌單
    candidates: dict[tuple, list[int]] = defaultdict(list)
    for pk, source, visual, title in Track.objects.values_list("pk", "song_source", "visual_source", "title"):
        candidates[(normalize(source), normalize(visual), normalize(title))].append(pk)

    by_date: dict[dt.date, list[PlayRow]] = defaultdict(list)
    for play in plays:
        by_date[play.date].append(play)

    festivals = {f.date: f for f in Festival.objects.filter(date__in=by_date.keys())}
    new_festivals = [
        Festival(date=date, name=guess_festival_name(date), kind="hall")
        for date in by_date if date not in festivals
    ]
    for festival in Festival.objects.bulk_create(new_festivals):
        festivals[festival.date] = festival

    # 每一場整批換新，重複匯入不會變成兩份
    Play.objects.filter(festival__date__in=by_date.keys()).delete()

    counts = {"matched": 0, "ambiguous": 0, "unmatched": 0}
    objs = []
    for date, rows in by_date.items():
        for position, row in enumerate(rows, start=1):
            matches = candidates.get((normalize(row.song_source), normalize(row.visual_source), normalize(row.title)), [])
            status = "matched" if len(matches) == 1 else "ambiguous" if matches else "unmatched"
            counts[status] += 1
            objs.append(Play(
                festival=festivals[date],
                position=position,
                song_source=row.song_source,
                visual_source=row.visual_source,
                title=row.title,
                song_key=song_key(row.song_source, row.title),
                track_id=matches[0] if status == "matched" else None,
                match_status=status,
            ))
    Play.objects.bulk_create(objs, batch_size=500)

    report.add(
        f"播放紀錄：{len(by_date)} 場（新建場次 {len(new_festivals)}）、{len(objs)} 首；"
        f"對到歌單 {counts['matched']}、有多個版本 {counts['ambiguous']}、歌單裡找不到 {counts['unmatched']}"
    )
    if counts["ambiguous"] or counts["unmatched"]:
        report.warn(
            "對不到的大多是沒有收進歌單的新曲、結尾的 5 秒畫面影片，屬於正常情況；"
            "它們一樣會計入播放次數（以歌曲出處 + 曲名計算），只是不會連到 DVD 編號。"
            "如果是打錯字（例如「甲鐵城」和「甲鉄城」），可以到後台「播放紀錄」用「對應狀態」篩選出來，修正試算表後再匯入一次。"
        )

    stale = Festival.objects.filter(plays__isnull=False).exclude(date__in=by_date.keys()).distinct()
    for festival in stale:
        report.warn(f"{festival} 在網站上有播放紀錄，但試算表裡沒有這個日期（沒有刪除）")


def _read(source, reader, report: ImportReport):
    workbook = load_workbook(source)
    try:
        return reader(workbook, report)
    finally:
        workbook.close()  # 唯讀模式會一直開著檔案，Windows 上不關會鎖住檔案


def import_songfest(songlist=None, history=None, *, prune: bool = False, dry_run: bool = False) -> ImportReport:
    """
    songlist：線上歌單、history：最新播放曲目統計（路徑或檔案物件，可以只給其中一個）。
    先匯入歌單再匯入播放紀錄，播放紀錄才能對到最新的歌單。dry_run 時全部復原，只回傳報告。
    """
    report = ImportReport()
    songs = _read(songlist, read_songlist, report) if songlist is not None else None
    plays = _read(history, read_play_history, report) if history is not None else None

    with transaction.atomic():
        if songs is not None:
            sync_tracks(songs, report, prune=prune)
        if plays is not None:
            sync_plays(plays, report)
        if dry_run:
            transaction.set_rollback(True)
            report.add("（試跑：以上變更都沒有寫入）")
    return report
